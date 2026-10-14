from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from typing import Literal

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from pydantic import BaseModel, Field

from backend.database.connection import get_db_connection

router = APIRouter(prefix="/api/exports", tags=["exports"])


class ExportRequest(BaseModel):
    ids: list[str] = Field(min_length=1, max_length=1000)
    currency: Literal["CNY", "USD", "KHR", "VND", "INR", "EUR", "HKD", "SGD"] = "CNY"


@router.post("/reimbursements.xlsx")
def export_reimbursements(request: ExportRequest):
    ids = list(dict.fromkeys(request.ids))
    with get_db_connection() as conn, conn.cursor() as cursor:
        cursor.execute('''SELECT r.id, r.date, d.name, r."desc", r.category,
                          r.amount, r.status, r.receipt_filename
                          FROM reimbursements r JOIN departments d ON d.id = r.department_id
                          WHERE r.id = ANY(%s)''', (ids,))
        found = {row[0]: row for row in cursor.fetchall()}
    if len(found) != len(ids):
        raise HTTPException(status_code=409, detail="部分明细已删除，请刷新列表后重新选择")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报销申请单"
    sheet.merge_cells("A1:H1")
    sheet["A1"] = "费用报销申请单"
    sheet["A1"].font = Font(size=20, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.row_dimensions[1].height = 36
    sheet.merge_cells("A2:H2")
    names = {"CNY": "人民币", "USD": "美元", "KHR": "柬埔寨瑞尔", "VND": "越南盾", "INR": "印度卢比", "EUR": "欧元", "HKD": "港币", "SGD": "新加坡元"}
    sheet["A2"] = f"申请人：________________    申请日期：{date.today().isoformat()}    币种：{names[request.currency]}（{request.currency}）"
    sheet.append(["序号", "日期", "部门", "报销事由", "类别", f"金额（{request.currency}）", "状态", "收据文件"])
    total = Decimal("0.00")
    for index, record_id in enumerate(ids, 1):
        _, day, department, description, category, amount, status, receipt = found[record_id]
        money = Decimal(str(amount)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        total += money
        sheet.append([index, day, department, description, category, money, status, receipt or ""])
        # User-entered strings must remain text, including values beginning with '='.
        for column in (2, 3, 4, 5, 7, 8):
            sheet.cell(sheet.max_row, column).data_type = "s"
        sheet.cell(sheet.max_row, 6).number_format = '#,##0.00'
    sheet.append(["合计", f"{len(ids)} 项", "", "", "", total, "", ""])
    sheet.cell(sheet.max_row, 6).number_format = '#,##0.00'
    last_table_row = sheet.max_row
    sheet.append([])
    sheet.append(["申请人签字：", "", "", "部门负责人：", "", "审批人："])
    sheet.merge_cells(start_row=sheet.max_row + 1, start_column=1, end_row=sheet.max_row + 1, end_column=8)
    sheet.cell(sheet.max_row, 1, "审批意见：________________________________________________________________")
    edge = Side(style="thin", color="CBD5E1")
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    for row in sheet.iter_rows(min_row=3, max_row=last_table_row, max_col=8):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row in (3, last_table_row):
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="E8EFF8")
        sheet.row_dimensions[row[0].row].height = 36
    for column, width in zip("ABCDEFGH", [8, 14, 16, 38, 14, 16, 12, 24]):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A4"
    sheet.print_title_rows = "1:3"
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=.3, right=.3, top=.5, bottom=.5)
    sheet.print_area = sheet.dimensions
    sheet.oddFooter.center.text = "第 &P 页 / 共 &N 页"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="reimbursement-application.xlsx"'})
